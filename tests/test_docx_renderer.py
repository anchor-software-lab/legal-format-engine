"""Tests for DOCX renderer."""

import tempfile
from pathlib import Path

from docx import Document as DocxDocument

from legal_format_engine.engines.pipeline import format_document
from legal_format_engine.renderers.docx_renderer import render_docx


class TestDocxRenderer:
    def test_creates_file(self, sample_brief_text, sample_metadata, wi_appellate_ruleset):
        doc = format_document(sample_brief_text, sample_metadata, wi_appellate_ruleset)
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            path = render_docx(doc, wi_appellate_ruleset, f.name)
        assert Path(path).exists()
        assert Path(path).stat().st_size > 0

    def test_valid_docx(self, sample_brief_text, sample_metadata, wi_appellate_ruleset):
        doc = format_document(sample_brief_text, sample_metadata, wi_appellate_ruleset)
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            path = render_docx(doc, wi_appellate_ruleset, f.name)
        # Should be readable by python-docx
        docx = DocxDocument(str(path))
        assert len(docx.paragraphs) > 0

    def test_margins(self, sample_brief_text, sample_metadata, wi_appellate_ruleset):
        doc = format_document(sample_brief_text, sample_metadata, wi_appellate_ruleset)
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            path = render_docx(doc, wi_appellate_ruleset, f.name)
        docx = DocxDocument(str(path))
        section = docx.sections[0]
        # 1.25 inch = 1143000 EMU, 2.0 inch = 1828800 EMU
        assert section.top_margin == 1143000
        assert section.left_margin == 1828800

    def test_contains_caption(self, sample_brief_text, sample_metadata, wi_appellate_ruleset):
        doc = format_document(sample_brief_text, sample_metadata, wi_appellate_ruleset)
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            path = render_docx(doc, wi_appellate_ruleset, f.name)
        docx = DocxDocument(str(path))
        all_text = " ".join(p.text for p in docx.paragraphs)
        assert "COURT OF APPEALS" in all_text
        assert "CHRISTOPHERSON" in all_text

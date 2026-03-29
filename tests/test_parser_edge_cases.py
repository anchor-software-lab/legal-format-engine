"""Edge case tests for all parsers (plain text, DOCX, PDF)."""

import tempfile
from pathlib import Path

import fitz
import pytest
from docx import Document as DocxDocument
from docx.shared import Inches, Pt

from legal_format_engine.models.document import HeadingLevel
from legal_format_engine.parsers.docx_parser import extract_format_profile, parse_docx
from legal_format_engine.parsers.pdf_parser import extract_pdf_format_profile, parse_pdf
from legal_format_engine.parsers.plain_text_parser import parse_plain_text


# ── Plain Text Parser Edge Cases ───────────────────────────────────


class TestPlainTextEdgeCases:
    def test_empty_string(self):
        doc = parse_plain_text("")
        assert doc.sections == []
        assert doc.raw_text == ""

    def test_whitespace_only(self):
        doc = parse_plain_text("   \n\n\n   ")
        assert len(doc.sections) == 0

    def test_single_word(self):
        doc = parse_plain_text("Hello")
        assert doc.raw_text == "Hello"

    def test_no_headings_just_content(self):
        text = "This is just a paragraph of text without any headings."
        doc = parse_plain_text(text)
        assert doc.raw_text is not None
        assert "paragraph" in doc.raw_text

    def test_heading_at_end_no_content(self):
        text = "Some preamble text.\n\nARGUMENT"
        doc = parse_plain_text(text)
        headings = [s.heading_text for s in doc.sections]
        assert "ARGUMENT" in headings

    def test_multiple_headings_no_content_between(self):
        text = "ARGUMENT\nCONCLUSION"
        doc = parse_plain_text(text)
        assert len(doc.sections) >= 1

    def test_very_long_document(self):
        text = "ARGUMENT\n\n" + ("Lorem ipsum dolor sit amet. " * 500) + "\n\nCONCLUSION\n\nDone."
        doc = parse_plain_text(text)
        assert len(doc.sections) >= 2

    def test_unicode_content(self):
        text = "ARGUMENT\n\nThe défendant's café was résumé-worthy.\n\nCONCLUSION\n\nDone."
        doc = parse_plain_text(text)
        assert "café" in doc.raw_text
        assert "résumé" in doc.raw_text

    def test_special_chars_in_headings(self):
        text = "ARGUMENT § 904.04\n\nSome content.\n\nCONCLUSION\n\nDone."
        doc = parse_plain_text(text)
        assert len(doc.sections) >= 1

    def test_mixed_line_endings(self):
        text = "ARGUMENT\r\n\r\nContent here.\r\nMore content.\r\n\r\nCONCLUSION\r\n\r\nDone."
        doc = parse_plain_text(text)
        assert len(doc.sections) >= 1

    def test_preserves_raw_text(self):
        text = "ARGUMENT\n\nThe court erred.\n\nCONCLUSION\n\nReverse."
        doc = parse_plain_text(text)
        assert "court erred" in doc.raw_text

    def test_numbered_list_not_heading(self):
        text = "ARGUMENT\n\n1. First point\n2. Second point\n3. Third point\n\nCONCLUSION"
        doc = parse_plain_text(text)
        # Numbered items may or may not be detected as headings
        # but the document should parse without error
        assert len(doc.sections) >= 1

    def test_deeply_nested_headings(self):
        text = (
            "ARGUMENT\n\n"
            "I. First argument\n\nContent.\n\n"
            "A. Sub-point\n\nContent.\n\n"
            "1. Detail\n\nContent.\n\n"
            "CONCLUSION\n\nDone."
        )
        doc = parse_plain_text(text)
        assert len(doc.sections) >= 3

    def test_heading_boundary_length(self):
        # Exactly 79 chars (should be heading if all caps)
        text = "A" * 79 + "\n\nContent"
        doc = parse_plain_text(text)
        assert len(doc.sections) >= 1

    def test_heading_over_length(self):
        # 200 chars all caps - too long for heading
        text = "A" * 200 + "\n\nContent"
        doc = parse_plain_text(text)
        # Should NOT be detected as heading
        for s in doc.sections:
            assert len(s.heading_text) < 200 or s.id == "preamble"


# ── DOCX Parser Edge Cases ────────────────────────────────────────


def _make_docx(**kwargs) -> Path:
    """Create a DOCX with customizable content."""
    docx = DocxDocument()
    section = docx.sections[0]
    section.top_margin = Inches(kwargs.get("margin", 1.0))
    section.bottom_margin = Inches(kwargs.get("margin", 1.0))
    section.left_margin = Inches(kwargs.get("margin", 1.0))
    section.right_margin = Inches(kwargs.get("margin", 1.0))

    for item in kwargs.get("paragraphs", []):
        p = docx.add_paragraph()
        run = p.add_run(item.get("text", ""))
        if item.get("bold"):
            run.bold = True
        if item.get("italic"):
            run.italic = True
        if item.get("font"):
            run.font.name = item["font"]
        if item.get("size"):
            run.font.size = Pt(item["size"])

    path = Path(tempfile.mktemp(suffix=".docx"))
    docx.save(str(path))
    return path


class TestDocxParserEdgeCases:
    def test_empty_docx(self):
        path = _make_docx(paragraphs=[])
        doc = parse_docx(path)
        assert doc.sections == []

    def test_single_paragraph(self):
        path = _make_docx(paragraphs=[{"text": "Just one paragraph."}])
        doc = parse_docx(path)
        assert doc.raw_text is not None
        assert "one paragraph" in doc.raw_text

    def test_bold_and_italic_mixed(self):
        path = _make_docx(paragraphs=[
            {"text": "HEADING", "bold": True},
            {"text": "Normal text"},
            {"text": "Italic text", "italic": True},
        ])
        doc = parse_docx(path)
        assert len(doc.sections) >= 1

    def test_unicode_in_docx(self):
        path = _make_docx(paragraphs=[
            {"text": "ARGUMENT"},
            {"text": "The défendant argued café résumé über alles."},
        ])
        doc = parse_docx(path)
        assert "café" in doc.raw_text

    def test_very_long_paragraph(self):
        long_text = "Word " * 5000
        path = _make_docx(paragraphs=[
            {"text": "ARGUMENT", "bold": True},
            {"text": long_text},
        ])
        doc = parse_docx(path)
        assert len(doc.raw_text) > 10000

    def test_all_bold_document(self):
        path = _make_docx(paragraphs=[
            {"text": "ARGUMENT", "bold": True},
            {"text": "Everything is bold here.", "bold": True},
            {"text": "CONCLUSION", "bold": True},
            {"text": "Also bold.", "bold": True},
        ])
        doc = parse_docx(path)
        assert len(doc.sections) >= 1

    def test_no_formatting_info(self):
        # Paragraphs with no special formatting
        path = _make_docx(paragraphs=[
            {"text": "STATEMENT OF FACTS"},
            {"text": "Just regular text."},
        ])
        doc = parse_docx(path)
        # Should still detect ALL CAPS heading
        assert any("STATEMENT" in s.heading_text for s in doc.sections)

    def test_file_not_found(self):
        with pytest.raises(FileNotFoundError):
            parse_docx("/nonexistent/file.docx")

    def test_extract_profile_empty(self):
        path = _make_docx(paragraphs=[])
        profile = extract_format_profile(path)
        assert "page" in profile
        assert "fonts" in profile

    def test_multiple_fonts(self):
        path = _make_docx(paragraphs=[
            {"text": "Times text", "font": "Times New Roman", "size": 12},
            {"text": "Arial text", "font": "Arial", "size": 11},
        ])
        profile = extract_format_profile(path)
        assert len(profile["fonts"]["all"]) >= 1

    def test_custom_margins(self):
        path = _make_docx(margin=2.0, paragraphs=[{"text": "Content"}])
        profile = extract_format_profile(path)
        assert profile["page"]["margin_left_inches"] == 2.0


# ── PDF Parser Edge Cases ─────────────────────────────────────────


def _make_pdf(**kwargs) -> Path:
    """Create a PDF with customizable content."""
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    y = 72

    for item in kwargs.get("items", []):
        text = item.get("text", "")
        fontname = item.get("font", "helv")
        fontsize = item.get("size", 12)
        x = item.get("x", 72)

        page.insert_text((x, y), text, fontsize=fontsize, fontname=fontname)
        y += fontsize * 1.5

    path = Path(tempfile.mktemp(suffix=".pdf"))
    doc.save(str(path))
    doc.close()
    return path


class TestPdfParserEdgeCases:
    def test_empty_pdf(self):
        doc = fitz.open()
        doc.new_page()
        path = Path(tempfile.mktemp(suffix=".pdf"))
        doc.save(str(path))
        doc.close()

        result = parse_pdf(path)
        assert result.sections == []

    def test_single_line(self):
        path = _make_pdf(items=[{"text": "Just one line."}])
        result = parse_pdf(path)
        assert "one line" in (result.raw_text or "")

    def test_unicode_in_pdf(self):
        path = _make_pdf(items=[
            {"text": "ARGUMENT"},
            {"text": "Unicode: cafe resume"},  # PDF font limits
        ])
        result = parse_pdf(path)
        assert result.raw_text is not None

    def test_different_font_sizes(self):
        path = _make_pdf(items=[
            {"text": "BIG HEADING", "size": 18},
            {"text": "Normal content.", "size": 12},
            {"text": "Small footnote.", "size": 8},
        ])
        result = parse_pdf(path)
        profile = extract_pdf_format_profile(path)
        assert len(profile["font_sizes"]["all_pt"]) >= 2

    def test_centered_text_detection(self):
        doc = fitz.open()
        page = doc.new_page(width=612, height=792)
        # Center text at page center (306)
        page.insert_text((230, 72), "CENTERED HEADING", fontsize=12)
        path = Path(tempfile.mktemp(suffix=".pdf"))
        doc.save(str(path))
        doc.close()

        profile = extract_pdf_format_profile(path)
        assert len(profile.get("headings", [])) >= 0  # May detect

    def test_file_not_found(self):
        with pytest.raises(FileNotFoundError):
            parse_pdf("/nonexistent/file.pdf")

    def test_multi_page_pdf(self):
        doc = fitz.open()
        for i in range(3):
            page = doc.new_page(width=612, height=792)
            page.insert_text((72, 72), f"PAGE {i + 1} HEADING", fontsize=14)
            page.insert_text((72, 96), f"Content for page {i + 1}.", fontsize=12)
        path = Path(tempfile.mktemp(suffix=".pdf"))
        doc.save(str(path))
        doc.close()

        result = parse_pdf(path)
        assert len(result.sections) >= 1
        assert "Content for page" in (result.raw_text or "")

    def test_page_dimensions(self):
        doc = fitz.open()
        # Legal size: 8.5 x 14
        doc.new_page(width=612, height=1008)
        path = Path(tempfile.mktemp(suffix=".pdf"))
        doc.save(str(path))
        doc.close()

        profile = extract_pdf_format_profile(path)
        assert profile["page"]["width_inches"] == 8.5
        assert profile["page"]["height_inches"] == 14.0

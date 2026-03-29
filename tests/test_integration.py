"""Integration tests: end-to-end pipeline, renderer, and real brief tests."""

import json
import tempfile
from pathlib import Path

import pytest
from click.testing import CliRunner
from docx import Document as DocxDocument

from legal_format_engine.cli import main
from legal_format_engine.engines.pipeline import (
    format_and_render_docx,
    format_and_render_markdown,
    format_document,
)
from legal_format_engine.models.metadata import (
    AttorneyInfo,
    CaseMetadata,
    DocumentMetadata,
    Party,
    PartyRole,
)
from legal_format_engine.parsers.docx_parser import parse_docx
from legal_format_engine.renderers.docx_renderer import render_docx
from legal_format_engine.renderers.markdown_renderer import render_markdown
from legal_format_engine.rules.loader import load_ruleset


def _make_meta():
    return DocumentMetadata(
        jurisdiction="wisconsin",
        court_level="appellate",
        document_type="brief",
        case=CaseMetadata(
            case_number="2025AP999999-CR",
            court_name="Court of Appeals",
            district="District II",
            county_of_origin="Dane",
            judge_name="Hon. Smith",
            parties=[
                Party(name="State of Wisconsin", role=PartyRole.PLAINTIFF_RESPONDENT),
                Party(name="John Test", role=PartyRole.DEFENDANT_APPELLANT),
            ],
        ),
        attorney=AttorneyInfo(
            name="Test Attorney",
            bar_number="9999999",
            firm="Test Firm",
            address="123 Test St",
            phone="(608) 555-0000",
            email="test@test.com",
        ),
        document_title="Brief of Defendant-Appellant",
    )


def _write_meta_json(meta):
    path = Path(tempfile.mktemp(suffix=".json"))
    path.write_text(json.dumps(meta.model_dump(mode="json")), encoding="utf-8")
    return path


# ── End-to-End Pipeline Tests ──────────────────────────────────────


class TestPipelineEndToEnd:
    def test_empty_text(self):
        meta = _make_meta()
        ruleset = load_ruleset("wisconsin", "appellate", "brief")
        doc = format_document("", meta, ruleset)
        # Should produce a document with generated stubs
        assert len(doc.sections) > 0
        assert len(doc.issues) > 0

    def test_whitespace_only_text(self):
        meta = _make_meta()
        ruleset = load_ruleset("wisconsin", "appellate", "brief")
        doc = format_document("   \n\n\n   ", meta, ruleset)
        assert isinstance(doc.sections, list)

    def test_minimal_valid_brief(self):
        text = (
            "STATEMENT OF ISSUES\n\n"
            "Whether the court erred.\n\n"
            "STATEMENT OF THE CASE\n\n"
            "Defendant was convicted.\n\n"
            "STATEMENT OF FACTS\n\n"
            "The facts are as follows.\n\n"
            "ARGUMENT\n\n"
            "The court erred because the evidence was insufficient.\n\n"
            "CONCLUSION\n\n"
            "For these reasons, reverse.\n"
        )
        meta = _make_meta()
        ruleset = load_ruleset("wisconsin", "appellate", "brief")
        doc = format_document(text, meta, ruleset, insert_missing=False)
        # Should find most required sections
        missing = [i for i in doc.issues if i.code == "MISSING_SECTION"]
        # Key sections present, but some optional ones missing
        found_codes = [i.section_id for i in doc.issues if i.code == "MISSING_SECTION"]
        assert "argument" not in found_codes
        assert "conclusion" not in found_codes

    def test_pipeline_no_insert(self):
        meta = _make_meta()
        ruleset = load_ruleset("wisconsin", "appellate", "brief")
        doc = format_document("ARGUMENT\n\nSome text.\n\nCONCLUSION\n\nDone.",
                              meta, ruleset, insert_missing=False)
        generated = [s for s in doc.sections if s.is_generated]
        assert len(generated) == 0

    def test_pipeline_with_insert(self):
        meta = _make_meta()
        ruleset = load_ruleset("wisconsin", "appellate", "brief")
        doc = format_document("ARGUMENT\n\nSome text.",
                              meta, ruleset, insert_missing=True)
        generated = [s for s in doc.sections if s.is_generated]
        assert len(generated) > 0

    def test_format_and_render_markdown(self):
        meta = _make_meta()
        ruleset = load_ruleset("wisconsin", "appellate", "brief")
        text = "ARGUMENT\n\nThe court erred.\n\nCONCLUSION\n\nReverse."
        doc, md = format_and_render_markdown(text, meta, ruleset)
        assert isinstance(md, str)
        assert "ARGUMENT" in md
        assert len(md) > 50

    def test_format_and_render_docx(self):
        meta = _make_meta()
        ruleset = load_ruleset("wisconsin", "appellate", "brief")
        text = "ARGUMENT\n\nThe court erred.\n\nCONCLUSION\n\nReverse."
        output = Path(tempfile.mktemp(suffix=".docx"))
        doc, path = format_and_render_docx(text, meta, ruleset, str(output))
        assert Path(path).exists()
        assert Path(path).stat().st_size > 0
        # Verify it's a valid DOCX
        docx = DocxDocument(str(path))
        all_text = "\n".join(p.text for p in docx.paragraphs)
        assert "ARGUMENT" in all_text

    def test_spd_variant_pipeline(self):
        meta = _make_meta()
        ruleset = load_ruleset("wisconsin", "appellate", "brief", variant="spd")
        text = (
            "ISSUES PRESENTED\n\n"
            "Whether the court erred.\n\n"
            "POSITION ON ORAL ARGUMENT AND PUBLICATION\n\n"
            "Not requested.\n\n"
            "STATEMENT OF THE CASE AND FACTS\n\n"
            "Defendant was convicted.\n\n"
            "ARGUMENT\n\n"
            "The court erred.\n\n"
            "CONCLUSION\n\n"
            "Reverse.\n"
        )
        doc = format_document(text, meta, ruleset, insert_missing=False)
        # SPD uses 10pt, 1.25" margins - should apply without error
        assert isinstance(doc.sections, list)


# ── Renderer Edge Cases ────────────────────────────────────────────


class TestRendererEdgeCases:
    def test_markdown_empty_document(self):
        from legal_format_engine.models.document import LegalDocument
        doc = LegalDocument(sections=[])
        md = render_markdown(doc)
        assert isinstance(md, str)

    def test_markdown_deep_heading_levels(self):
        from legal_format_engine.models.document import (
            ContentBlock, HeadingLevel, LegalDocument, Section,
        )
        doc = LegalDocument(sections=[
            Section(id="l1", heading_text="Level 1",
                    heading_level=HeadingLevel.LEVEL_1,
                    content=[ContentBlock(text="Content")]),
            Section(id="l2", heading_text="Level 2",
                    heading_level=HeadingLevel.LEVEL_2,
                    content=[ContentBlock(text="Content")]),
            Section(id="l3", heading_text="Level 3",
                    heading_level=HeadingLevel.LEVEL_3,
                    content=[ContentBlock(text="Content")]),
            Section(id="l4", heading_text="Level 4",
                    heading_level=HeadingLevel.LEVEL_4,
                    content=[ContentBlock(text="Content")]),
        ])
        md = render_markdown(doc)
        assert "# Level 1" in md
        assert "## Level 2" in md

    def test_docx_unicode_content(self):
        from legal_format_engine.models.document import (
            CaptionBlock, ContentBlock, HeadingLevel,
            LegalDocument, Section,
        )
        doc = LegalDocument(
            caption=CaptionBlock(lines=[ContentBlock(text="COURT")]),
            sections=[
                Section(id="arg", heading_text="ARGUMENT",
                        heading_level=HeadingLevel.LEVEL_1,
                        content=[ContentBlock(text="Über café résumé naïve")]),
            ],
        )
        ruleset = load_ruleset("wisconsin", "appellate", "brief")
        output = Path(tempfile.mktemp(suffix=".docx"))
        render_docx(doc, ruleset, str(output))
        assert output.exists()


# ── CLI Edge Cases ─────────────────────────────────────────────────


class TestCLIEdgeCases:
    def test_format_invalid_metadata(self):
        runner = CliRunner()
        with runner.isolated_filesystem():
            Path("input.txt").write_text("ARGUMENT\n\nContent.")
            Path("meta.json").write_text("not valid json")
            result = runner.invoke(main, [
                "format", "input.txt", "-o", "out.docx", "-m", "meta.json",
            ])
            assert result.exit_code != 0

    def test_format_missing_input(self):
        runner = CliRunner()
        result = runner.invoke(main, [
            "format", "nonexistent.txt", "-o", "out.docx", "-m", "meta.json",
        ])
        assert result.exit_code != 0

    def test_format_with_variant(self):
        runner = CliRunner()
        meta = _make_meta()
        with runner.isolated_filesystem():
            Path("input.txt").write_text("ARGUMENT\n\nContent.\n\nCONCLUSION\n\nDone.")
            meta_path = _write_meta_json(meta)
            result = runner.invoke(main, [
                "format", "input.txt", "-o", "out.docx", "-m", str(meta_path),
                "--variant", "spd",
            ])
            assert result.exit_code == 0
            assert Path("out.docx").exists()

    def test_validate_clean_doc(self):
        runner = CliRunner()
        meta = _make_meta()
        text = (
            "STATEMENT OF ISSUES\n\nIssue.\n\n"
            "STATEMENT OF THE CASE\n\nCase.\n\n"
            "STATEMENT OF FACTS\n\nFacts.\n\n"
            "ARGUMENT\n\nArgument.\n\n"
            "CONCLUSION\n\nDone.\n"
        )
        with runner.isolated_filesystem():
            Path("input.txt").write_text(text)
            meta_path = _write_meta_json(meta)
            result = runner.invoke(main, [
                "validate", "-i", "input.txt", "-m", str(meta_path),
            ])
            # May have warnings for missing certifications but should not crash
            assert result.exit_code in (0, 1)

    def test_citations_command(self):
        runner = CliRunner()
        text = "State v. Sullivan, 216 Wis. 2d 768 (1998). Id. at 774."
        with runner.isolated_filesystem():
            Path("brief.txt").write_text(text)
            result = runner.invoke(main, ["citations", "brief.txt"])
            assert result.exit_code == 0
            assert "citation" in result.output.lower()

    def test_analyze_nonexistent(self):
        runner = CliRunner()
        result = runner.invoke(main, ["analyze", "/nonexistent/file.txt"])
        assert result.exit_code != 0


# ── Real Brief Test ────────────────────────────────────────────────


class TestRealChristophersonBrief:
    """Integration tests against the actual Christopherson brief."""

    @pytest.fixture
    def brief_path(self):
        path = Path("reference/briefs/Christopherson_BIC_final.docx")
        if not path.exists():
            pytest.skip("Christopherson brief not available")
        return path

    def test_parses_without_error(self, brief_path):
        doc = parse_docx(brief_path)
        assert len(doc.sections) > 10
        assert doc.raw_text is not None
        assert len(doc.raw_text) > 1000

    def test_detects_key_sections(self, brief_path):
        doc = parse_docx(brief_path)
        headings = [s.heading_text.upper() for s in doc.sections]
        heading_text = " ".join(headings)
        assert "TABLE OF CONTENTS" in heading_text
        assert "ARGUMENT" in heading_text

    def test_full_pipeline_spd(self, brief_path):
        doc = parse_docx(brief_path)
        text = doc.raw_text or ""
        meta = DocumentMetadata(
            jurisdiction="wisconsin",
            court_level="appellate",
            document_type="brief",
            case=CaseMetadata(
                case_number="2025AP002377-CR",
                court_name="Court of Appeals",
                district="District II",
                county_of_origin="Calumet",
                judge_name="Hon. Carey Reed",
                parties=[
                    Party(name="State of Wisconsin", role=PartyRole.PLAINTIFF_RESPONDENT),
                    Party(name="Jeffrey James Christopherson", role=PartyRole.DEFENDANT_APPELLANT),
                ],
            ),
            attorney=AttorneyInfo(
                name="Nicholas G. Smith",
                bar_number="1089586",
                firm="Office of the State Public Defender",
                address="Post Office Box 7862\nMadison, WI 53707-7862",
                phone="(608) 261-5417",
                email="smithn@opd.wi.gov",
            ),
            document_title="Brief of Defendant-Appellant",
        )
        ruleset = load_ruleset("wisconsin", "appellate", "brief", variant="spd")
        result = format_document(text, meta, ruleset, insert_missing=False)
        assert len(result.sections) > 0

    def test_citation_check_on_brief(self, brief_path):
        from legal_format_engine.engines.citation_engine import check_citation_consistency

        doc = parse_docx(brief_path)
        text = doc.raw_text or ""
        report = check_citation_consistency(text)
        # Brief should have many citations
        assert len(report.citations) > 20
        assert len(report.statute_citations) > 5

    def test_docx_roundtrip(self, brief_path):
        """Parse the brief, format it, render to DOCX, parse again."""
        doc = parse_docx(brief_path)
        text = doc.raw_text or ""
        meta = DocumentMetadata(
            jurisdiction="wisconsin",
            court_level="appellate",
            document_type="brief",
            case=CaseMetadata(
                case_number="2025AP002377-CR",
                court_name="Court of Appeals",
                parties=[
                    Party(name="State of Wisconsin", role=PartyRole.PLAINTIFF_RESPONDENT),
                    Party(name="Jeffrey James Christopherson", role=PartyRole.DEFENDANT_APPELLANT),
                ],
            ),
            attorney=AttorneyInfo(
                name="Nicholas G. Smith", bar_number="1089586",
                address="PO Box 7862", phone="(608) 261-5417",
                email="smithn@opd.wi.gov",
            ),
            document_title="Brief of Defendant-Appellant",
        )
        ruleset = load_ruleset("wisconsin", "appellate", "brief")
        output = Path(tempfile.mktemp(suffix=".docx"))
        formatted, path = format_and_render_docx(
            text, meta, ruleset, str(output), word_count=15000
        )
        # Re-parse the output
        reparsed = parse_docx(path)
        assert len(reparsed.sections) > 0
        assert reparsed.raw_text is not None

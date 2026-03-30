"""Tests for DOCX parser."""

import tempfile
from pathlib import Path

from docx import Document as DocxDocument

from legal_format_engine.parsers.docx_parser import extract_format_profile, parse_docx


def _create_sample_docx() -> Path:
    """Create a minimal DOCX for testing."""
    docx = DocxDocument()

    # Set margins
    section = docx.sections[0]
    from docx.shared import Inches
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    # Add content
    p = docx.add_paragraph()
    run = p.add_run("STATEMENT OF ISSUES")
    run.bold = True
    from docx.shared import Pt
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    docx.add_paragraph("Whether the court erred in its ruling.")

    p2 = docx.add_paragraph()
    run2 = p2.add_run("ARGUMENT")
    run2.bold = True
    run2.font.name = "Times New Roman"
    run2.font.size = Pt(12)

    docx.add_paragraph("The court clearly erred because...")

    p3 = docx.add_paragraph()
    run3 = p3.add_run("CONCLUSION")
    run3.bold = True
    run3.font.name = "Times New Roman"
    run3.font.size = Pt(12)

    docx.add_paragraph("For these reasons, relief should be granted.")

    path = Path(tempfile.mktemp(suffix=".docx"))
    docx.save(str(path))
    return path


class TestDocxParser:
    def test_parse_basic(self):
        path = _create_sample_docx()
        doc = parse_docx(path)
        assert len(doc.sections) > 0
        assert doc.raw_text is not None

    def test_detects_headings(self):
        path = _create_sample_docx()
        doc = parse_docx(path)
        headings = [s.heading_text for s in doc.sections]
        assert "STATEMENT OF ISSUES" in headings
        assert "ARGUMENT" in headings
        assert "CONCLUSION" in headings

    def test_preserves_content(self):
        path = _create_sample_docx()
        doc = parse_docx(path)
        arg = next(s for s in doc.sections if s.heading_text == "ARGUMENT")
        all_text = " ".join(b.text for b in arg.content)
        assert "clearly erred" in all_text

    def test_file_not_found(self):
        import pytest
        with pytest.raises(FileNotFoundError):
            parse_docx("/nonexistent/file.docx")


class TestFormatProfile:
    def test_extracts_profile(self):
        path = _create_sample_docx()
        profile = extract_format_profile(path)
        assert "page" in profile
        assert "fonts" in profile
        assert "headings" in profile

    def test_detects_margins(self):
        path = _create_sample_docx()
        profile = extract_format_profile(path)
        assert profile["page"]["margin_top_inches"] == 1.0
        assert profile["page"]["margin_left_inches"] == 1.0

    def test_detects_font(self):
        path = _create_sample_docx()
        profile = extract_format_profile(path)
        assert profile["fonts"]["dominant"] == "Times New Roman"

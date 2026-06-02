"""Shared docx fixture builders for tests.

Importable from any package's conftest so we don't duplicate the same
40-line "synthesize a brief" code in every package. Generating the
fixture at test time (rather than checking in a .docx binary) keeps the
source tree text-only and makes the fixture's shape auditable.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt


def make_sample_brief(out_path: str | Path) -> Path:
    """Build a 4-paragraph brief and write it to `out_path`.

    Shape: heading, body paragraph (intentionally 12pt to trigger a
    finding when the rule expects 13pt), block quote, conclusion. The
    body paragraph contains real citations (Tews, Brown) so the citation
    checkers have something to chew on too.
    """
    out_path = Path(out_path)
    doc = DocxDocument()

    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    heading = doc.add_heading("ARGUMENT I. STANDARD OF REVIEW", level=1)
    heading.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

    body = doc.add_paragraph()
    body.paragraph_format.first_line_indent = Inches(0.5)
    body.paragraph_format.line_spacing = 2.0
    body_run = body.add_run(
        "This Court reviews a circuit court's grant of summary judgment de "
        "novo. See Tews v. NHI, LLC, 2010 WI 137, ¶ 4, 330 Wis. 2d 389. "
        "See also Brown v. Holiday, 2008 WI 49."
    )
    body_run.font.name = "Times New Roman"
    body_run.font.size = Pt(12)

    quote = doc.add_paragraph()
    quote.paragraph_format.left_indent = Inches(0.5)
    quote.paragraph_format.line_spacing = 1.0
    quote_run = quote.add_run(
        "Summary judgment is appropriate if the pleadings, depositions, "
        "answers to interrogatories, and admissions on file show that "
        "there is no genuine issue as to any material fact."
    )
    quote_run.font.name = "Times New Roman"
    quote_run.font.size = Pt(13)

    concl = doc.add_paragraph()
    concl.paragraph_format.first_line_indent = Inches(0.5)
    concl.paragraph_format.line_spacing = 2.0
    concl_run = concl.add_run(
        "For the foregoing reasons, the order of the circuit court should "
        "be affirmed."
    )
    concl_run.font.name = "Times New Roman"
    concl_run.font.size = Pt(13)

    doc.save(str(out_path))
    return out_path


def make_well_formatted_brief(out_path: str | Path) -> Path:
    """A 3-paragraph brief that should produce zero findings against
    the standard 13pt / 2.0-spacing / 1" margin rules.

    Used as the negative case in the fixture corpus: any new checker
    that mistakenly flags clean text will fail this fixture.
    """
    out_path = Path(out_path)
    doc = DocxDocument()
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    h = doc.add_heading("ARGUMENT", level=1)
    h.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

    body = doc.add_paragraph()
    body.paragraph_format.first_line_indent = Inches(0.5)
    body.paragraph_format.line_spacing = 2.0
    body_run = body.add_run(
        "The standard of review is de novo. See Tews v. NHI, LLC, 2010 WI "
        "137, ¶ 4, 330 Wis. 2d 389. The trial court erred."
    )
    body_run.font.name = "Times New Roman"
    body_run.font.size = Pt(13)

    concl = doc.add_paragraph()
    concl.paragraph_format.first_line_indent = Inches(0.5)
    concl.paragraph_format.line_spacing = 2.0
    concl_run = concl.add_run("The order should be affirmed.")
    concl_run.font.name = "Times New Roman"
    concl_run.font.size = Pt(13)

    doc.save(str(out_path))
    return out_path


def make_brief_with_malformed_signal(out_path: str | Path) -> Path:
    """A brief whose body paragraph uses "See, also," instead of "See also".

    The Bluebook signal checker should flag BB.SIGNAL.UNKNOWN on the
    Brown cite.
    """
    out_path = Path(out_path)
    doc = DocxDocument()
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    h = doc.add_heading("ARGUMENT", level=1)
    h.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

    body = doc.add_paragraph()
    body.paragraph_format.first_line_indent = Inches(0.5)
    body.paragraph_format.line_spacing = 2.0
    body_run = body.add_run(
        "We review de novo. See Tews v. NHI, LLC, 2010 WI 137, ¶ 4. "
        "See, also, Brown v. Holiday, 2008 WI 49, ¶ 12."
    )
    body_run.font.name = "Times New Roman"
    body_run.font.size = Pt(13)

    doc.save(str(out_path))
    return out_path


def make_brief_with_orphan_id(out_path: str | Path) -> Path:
    """A brief that opens with `Id. at ¶ 5.` — Bluebook Rule 10.9 forbids
    Id. without a preceding citation; the short-form checker should flag
    BB.SHORT_FORM.ORPHAN_ID.
    """
    out_path = Path(out_path)
    doc = DocxDocument()
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    h = doc.add_heading("ARGUMENT", level=1)
    h.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

    body = doc.add_paragraph()
    body.paragraph_format.first_line_indent = Inches(0.5)
    body.paragraph_format.line_spacing = 2.0
    body_run = body.add_run(
        "The standard of review is well established. Id. at ¶ 5. "
        "See Tews v. NHI, LLC, 2010 WI 137, ¶ 4."
    )
    body_run.font.name = "Times New Roman"
    body_run.font.size = Pt(13)

    doc.save(str(out_path))
    return out_path


# Public corpus map. Tests parametrize over this; CLI demos use it to
# generate sample documents.
FIXTURE_BUILDERS: dict[str, callable] = {
    "well_formatted": make_well_formatted_brief,
    "font_size_violation": make_sample_brief,
    "malformed_signal": make_brief_with_malformed_signal,
    "orphan_id": make_brief_with_orphan_id,
}
